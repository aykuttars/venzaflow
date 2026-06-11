from __future__ import annotations

import csv
import io
from decimal import Decimal, InvalidOperation

from apps.products.models import Category, Product


def export_products_csv(tenant_id: int) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["sku", "barcode", "name", "category_slug", "unit_price", "cost_price", "is_active"])
    for p in Product.objects.filter(tenant_id=tenant_id).select_related("category"):
        writer.writerow(
            [
                p.sku,
                p.barcode,
                p.name,
                p.category.slug,
                p.unit_price,
                p.cost_price or "",
                "1" if p.is_active else "0",
            ]
        )
    return buf.getvalue()


def import_products_csv(tenant_id: int, raw: bytes) -> dict:
    text = raw.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    created = 0
    updated = 0
    errors: list[dict] = []

    for i, row in enumerate(reader, start=2):
        sku = (row.get("sku") or "").strip()
        if not sku:
            errors.append({"row": i, "error": "sku required"})
            continue
        slug = (row.get("category_slug") or row.get("category") or "default").strip()
        category, _ = Category.objects.get_or_create(
            tenant_id=tenant_id,
            slug=slug,
            defaults={"name": slug.replace("-", " ").title()},
        )
        try:
            unit_price = Decimal(str(row.get("unit_price") or "0"))
        except InvalidOperation:
            errors.append({"row": i, "error": "invalid unit_price"})
            continue
        cost_raw = row.get("cost_price")
        cost_price = None
        if cost_raw not in (None, ""):
            try:
                cost_price = Decimal(str(cost_raw))
            except InvalidOperation:
                errors.append({"row": i, "error": "invalid cost_price"})
                continue
        defaults = {
            "name": (row.get("name") or sku).strip(),
            "category": category,
            "unit_price": unit_price,
            "cost_price": cost_price,
            "barcode": (row.get("barcode") or "").strip(),
            "is_active": str(row.get("is_active", "1")).strip() not in ("0", "false", "False"),
        }
        obj, was_created = Product.objects.update_or_create(
            tenant_id=tenant_id,
            sku=sku,
            defaults=defaults,
        )
        if was_created:
            created += 1
        else:
            updated += 1

    return {"created": created, "updated": updated, "errors": errors}


def export_products_xlsx(tenant_id: int) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "products"
    ws.append(["sku", "barcode", "name", "category_slug", "unit_price", "cost_price", "is_active"])
    for p in Product.objects.filter(tenant_id=tenant_id).select_related("category"):
        ws.append(
            [
                p.sku,
                p.barcode,
                p.name,
                p.category.slug,
                float(p.unit_price),
                float(p.cost_price) if p.cost_price is not None else "",
                1 if p.is_active else 0,
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_products_xlsx(tenant_id: int, raw: bytes) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "updated": 0, "errors": [{"row": 1, "error": "empty file"}]}
    headers = [str(h or "").strip().lower() for h in rows[0]]
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows[1:]:
        writer.writerow(["" if c is None else c for c in row])
    return import_products_csv(tenant_id, buf.getvalue().encode("utf-8"))
