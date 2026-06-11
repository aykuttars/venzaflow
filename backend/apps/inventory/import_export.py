from __future__ import annotations

import csv
import io

from apps.inventory.models import Stock, Warehouse
from apps.products.models import Product


def export_stock_csv(tenant_id: int) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["product_sku", "warehouse_code", "location_code", "quantity", "reserved_quantity", "reorder_level"])
    qs = Stock.objects.filter(tenant_id=tenant_id).select_related("product", "warehouse", "location")
    for s in qs:
        writer.writerow(
            [
                s.product.sku,
                s.warehouse.code,
                s.location.code if s.location_id else "",
                s.quantity,
                s.reserved_quantity,
                s.reorder_level,
            ]
        )
    return buf.getvalue()


def import_stock_csv(tenant_id: int, raw: bytes) -> dict:
    text = raw.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    created = 0
    updated = 0
    errors: list[dict] = []

    for i, row in enumerate(reader, start=2):
        sku = (row.get("product_sku") or "").strip()
        wh_code = (row.get("warehouse_code") or "").strip()
        if not sku or not wh_code:
            errors.append({"row": i, "error": "product_sku and warehouse_code required"})
            continue
        try:
            product = Product.objects.get(tenant_id=tenant_id, sku=sku)
            warehouse = Warehouse.objects.get(tenant_id=tenant_id, code=wh_code)
        except (Product.DoesNotExist, Warehouse.DoesNotExist):
            errors.append({"row": i, "error": "product or warehouse not found"})
            continue
        location = None
        loc_code = (row.get("location_code") or "").strip()
        if loc_code:
            from apps.inventory.models import Location

            location = Location.objects.filter(
                tenant_id=tenant_id, warehouse=warehouse, code=loc_code
            ).first()
            if not location:
                errors.append({"row": i, "error": "location not found"})
                continue
        try:
            qty = int(row.get("quantity") or 0)
            reserved = int(row.get("reserved_quantity") or 0)
            reorder = int(row.get("reorder_level") or 0)
        except ValueError:
            errors.append({"row": i, "error": "invalid quantity"})
            continue
        _, was_created = Stock.objects.update_or_create(
            tenant_id=tenant_id,
            product=product,
            warehouse=warehouse,
            location=location,
            defaults={
                "quantity": qty,
                "reserved_quantity": reserved,
                "reorder_level": reorder,
            },
        )
        if was_created:
            created += 1
        else:
            updated += 1

    return {"created": created, "updated": updated, "errors": errors}


def export_stock_xlsx(tenant_id: int) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "stock"
    ws.append(["product_sku", "warehouse_code", "location_code", "quantity", "reserved_quantity", "reorder_level"])
    qs = Stock.objects.filter(tenant_id=tenant_id).select_related("product", "warehouse", "location")
    for s in qs:
        ws.append(
            [
                s.product.sku,
                s.warehouse.code,
                s.location.code if s.location_id else "",
                s.quantity,
                s.reserved_quantity,
                s.reorder_level,
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_stock_xlsx(tenant_id: int, raw: bytes) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "updated": 0, "errors": [{"row": 1, "error": "empty file"}]}
    headers = [str(h or "").strip().lower() for h in rows[0]]
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows[1:]:
        writer.writerow(["" if c is None else c for c in row])
    return import_stock_csv(tenant_id, buf.getvalue().encode("utf-8"))
